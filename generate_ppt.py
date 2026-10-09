import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def set_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def create_textbox(slide, left, top, width, height, text, font_size, font_color, bold=False, align=PP_ALIGN.LEFT, font_name='Segoe UI'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.font.name = font_name
    return txBox

def add_bullet(tf, text, font_size, font_color, level=0, bold=False, font_name='Segoe UI'):
    p = tf.add_paragraph()
    p.text = text
    p.level = level
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.font.name = font_name
    return p

def create_presentation():
    prs = Presentation()
    # 16:9 Aspect Ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_COLOR = RGBColor(12, 35, 64)         # Dark Navy
    TEXT_MAIN = RGBColor(255, 255, 255)     # White
    TEXT_SUB = RGBColor(200, 208, 217)      # Light Gray
    ACCENT = RGBColor(0, 210, 211)          # Cyan
    HIGHLIGHT = RGBColor(253, 160, 91)      # Orange/Gold

    blank_layout = prs.slide_layouts[6] # Blank layout

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide, BG_COLOR)

    # Decorative Line
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(3.5), Inches(13.333), Inches(0.1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.color.rgb = ACCENT

    create_textbox(slide, Inches(1), Inches(1.8), Inches(11), Inches(1.5), 
                   "Commercial Fleet Operations &\nTurnaround Optimization Model", 
                   44, TEXT_MAIN, bold=True, align=PP_ALIGN.CENTER)

    create_textbox(slide, Inches(1), Inches(4.0), Inches(11), Inches(1), 
                   "Predicting Cascading Delays to Maximize Aircraft Utilization and Reduce Operational Costs", 
                   24, TEXT_SUB, align=PP_ALIGN.CENTER)
                   
    create_textbox(slide, Inches(1), Inches(6.5), Inches(11), Inches(0.5), 
                   "End-to-End Analytics & Machine Learning Pipeline", 
                   16, ACCENT, bold=True, align=PP_ALIGN.CENTER)


    # =========================================================================
    # SLIDE 2: THE BUSINESS PROBLEM
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide, BG_COLOR)
    
    # Title
    create_textbox(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(1), 
                   "The Operational Bottleneck: Cascading Delays", 
                   36, ACCENT, bold=True)
    
    # Main Content Box
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    add_bullet(tf, "Airlines operate on incredibly tight turnaround schedules to maximize aircraft utilization.", 24, TEXT_MAIN, bold=True)
    
    add_bullet(tf, "The Turnaround Buffer", 22, HIGHLIGHT, level=1, bold=True)
    add_bullet(tf, "The scheduled ground time allocated for cleaning, fueling, and boarding between an aircraft's arrival and its next departure.", 18, TEXT_SUB, level=2)
    
    add_bullet(tf, "The Domino Effect (Cascading Delays)", 22, HIGHLIGHT, level=1, bold=True)
    add_bullet(tf, "When an aircraft arrives late, or ground operations exceed the buffer, the delay directly propagates to the next flight utilizing that same physical aircraft.", 18, TEXT_SUB, level=2)
    
    add_bullet(tf, "The Financial Hemorrhage", 22, HIGHLIGHT, level=1, bold=True)
    add_bullet(tf, "According to the FAA and Airlines for America, every single minute of delay costs an airline approximately $75 in crew overtime, passenger compensations, and fuel inefficiencies.", 18, TEXT_SUB, level=2)


    # =========================================================================
    # SLIDE 3: DATA ENGINEERING & ARCHITECTURE
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide, BG_COLOR)
    
    create_textbox(slide, Inches(0.8), Inches(0.5), Inches(10), Inches(1), 
                   "Phase 1: Robust Data Ingestion & Engineering", 
                   36, ACCENT, bold=True)
                   
    # Architecture Blocks
    # Block 1
    s1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.5), Inches(3.2), Inches(3))
    s1.fill.solid()
    s1.fill.fore_color.rgb = RGBColor(20, 50, 90)
    s1.line.color.rgb = ACCENT
    create_textbox(slide, Inches(1.1), Inches(2.7), Inches(3), Inches(0.5), "1. Raw Data Parsing", 22, HIGHLIGHT, bold=True, align=PP_ALIGN.CENTER)
    create_textbox(slide, Inches(1.1), Inches(3.5), Inches(3), Inches(1.5), 
                   "Scans root directories for CSV/ZIP datasets.\n\nCleans missing values and transforms complex float time-series data into standard Python DateTime objects.", 
                   14, TEXT_SUB, align=PP_ALIGN.CENTER)

    # Block 2
    s2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5), Inches(2.5), Inches(3.2), Inches(3))
    s2.fill.solid()
    s2.fill.fore_color.rgb = RGBColor(20, 50, 90)
    s2.line.color.rgb = ACCENT
    create_textbox(slide, Inches(5.1), Inches(2.7), Inches(3), Inches(0.5), "2. Transformation", 22, HIGHLIGHT, bold=True, align=PP_ALIGN.CENTER)
    create_textbox(slide, Inches(5.1), Inches(3.5), Inches(3), Inches(1.5), 
                   "Mathematically derives Actual vs Scheduled durations.\n\nHeuristically adjusts for complex edge-cases like overnight flight arrivals spanning multiple days.", 
                   14, TEXT_SUB, align=PP_ALIGN.CENTER)

    # Block 3
    s3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9), Inches(2.5), Inches(3.2), Inches(3))
    s3.fill.solid()
    s3.fill.fore_color.rgb = RGBColor(20, 50, 90)
    s3.line.color.rgb = ACCENT
    create_textbox(slide, Inches(9.1), Inches(2.7), Inches(3), Inches(0.5), "3. SQLite Database", 22, HIGHLIGHT, bold=True, align=PP_ALIGN.CENTER)
    create_textbox(slide, Inches(9.1), Inches(3.5), Inches(3), Inches(1.5), 
                   "Loads clean data into a local 'operations.db' database.\n\nBuilds B-Tree indexes on Tail Numbers and Origins to guarantee rapid downstream SQL querying.", 
                   14, TEXT_SUB, align=PP_ALIGN.CENTER)


    # =========================================================================
    # SLIDE 4: ANALYTICAL SQL ENGINE
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide, BG_COLOR)
    
    create_textbox(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(1), 
                   "Phase 2: The Analytical SQL Engine", 
                   36, ACCENT, bold=True)
                   
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    add_bullet(tf, "Unlocking Network Intelligence using Advanced Window Functions", 24, TEXT_MAIN, bold=True)
    
    add_bullet(tf, "Common Table Expressions (CTEs)", 22, HIGHLIGHT, level=1, bold=True)
    add_bullet(tf, "Instead of complex Python loops, we utilize highly optimized SQL CTEs to trace aircraft history.", 18, TEXT_SUB, level=2)
    
    add_bullet(tf, "Physical Aircraft Sequencing [ LAG() ]", 22, HIGHLIGHT, level=1, bold=True)
    add_bullet(tf, "By partitioning data by 'Tail_Number' and ordering by 'Scheduled_Departure', we dynamically lookup the exact arrival time and location of the aircraft's prior flight.", 18, TEXT_SUB, level=2)
    add_bullet(tf, "This confirms physical aircraft continuity (Prior Destination = Current Origin) to calculate true ground turnaround metrics.", 18, TEXT_SUB, level=2)
    
    add_bullet(tf, "Isolating Network Choke Points", 22, HIGHLIGHT, level=1, bold=True)
    add_bullet(tf, "We aggregate these metrics by Airport and Carrier, flagging flights where an inbound delay directly triggered an outbound delay, effectively creating a map of operational choke points.", 18, TEXT_SUB, level=2)


    # =========================================================================
    # SLIDE 5: PREDICTIVE MACHINE LEARNING
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide, BG_COLOR)
    
    create_textbox(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(1), 
                   "Phase 3: Machine Learning & Prediction", 
                   36, ACCENT, bold=True)
                   
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    add_bullet(tf, "The Goal", 24, HIGHLIGHT, bold=True)
    add_bullet(tf, "Predict if a flight will breach its turnaround buffer and cause severe downstream delays.", 18, TEXT_SUB, level=1)
    
    add_bullet(tf, "Feature Engineering", 24, HIGHLIGHT, bold=True)
    add_bullet(tf, "Turn Buffer Volume (Mins)", 16, TEXT_MAIN, level=1)
    add_bullet(tf, "Inbound Delay Stress", 16, TEXT_MAIN, level=1)
    add_bullet(tf, "Route Congestion Profiles", 16, TEXT_MAIN, level=1)
    add_bullet(tf, "Tail-Number Historical Delay Frequencies", 16, TEXT_MAIN, level=1)
    
    add_bullet(tf, "The Algorithm", 24, HIGHLIGHT, bold=True)
    add_bullet(tf, "Random Forest Classifier. We explicitly combat dataset class-imbalance by utilizing computationally balanced class weights.", 18, TEXT_SUB, level=1)

    # Metrics Box
    s_metrics = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(2), Inches(5), Inches(4))
    s_metrics.fill.solid()
    s_metrics.fill.fore_color.rgb = RGBColor(20, 50, 90)
    s_metrics.line.color.rgb = ACCENT
    
    create_textbox(slide, Inches(7.5), Inches(2.2), Inches(5), Inches(0.5), "Model Performance", 28, HIGHLIGHT, bold=True, align=PP_ALIGN.CENTER)
    create_textbox(slide, Inches(7.8), Inches(3.0), Inches(4.4), Inches(2.5), 
                   "PR-AUC Score: 0.82\n\nPrecision: 78.4%\n(High accuracy when flagging delays)\n\nRecall: 65.9%\n(Caught 66% of all actual delays)", 
                   20, TEXT_MAIN, align=PP_ALIGN.LEFT)


    # =========================================================================
    # SLIDE 6: FINANCIAL ROI
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_background(slide, BG_COLOR)
    
    create_textbox(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(1), 
                   "Phase 4: Operational Financial Translation", 
                   36, ACCENT, bold=True)
                   
    create_textbox(slide, Inches(0.8), Inches(1.5), Inches(11), Inches(0.8), 
                   "Translating data science metrics into executive financial value.", 
                   24, TEXT_SUB)

    # Box 1: Exposure
    s_exp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(2.8), Inches(4.5), Inches(2.5))
    s_exp.fill.solid()
    s_exp.fill.fore_color.rgb = RGBColor(40, 20, 20)  # Dark Red tint
    s_exp.line.color.rgb = RGBColor(255, 50, 50)
    
    create_textbox(slide, Inches(1.5), Inches(3.2), Inches(4.5), Inches(0.5), "Predicted Delay Cost Exposure", 22, TEXT_SUB, align=PP_ALIGN.CENTER)
    create_textbox(slide, Inches(1.5), Inches(3.8), Inches(4.5), Inches(1.0), "$51.2 Million", 44, RGBColor(255, 100, 100), bold=True, align=PP_ALIGN.CENTER)

    # Box 2: Savings
    s_sav = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.3), Inches(2.8), Inches(4.5), Inches(2.5))
    s_sav.fill.solid()
    s_sav.fill.fore_color.rgb = RGBColor(20, 40, 20)  # Dark Green tint
    s_sav.line.color.rgb = RGBColor(50, 255, 50)
    
    create_textbox(slide, Inches(7.3), Inches(3.2), Inches(4.5), Inches(0.5), "Potential Operational Savings", 22, TEXT_SUB, align=PP_ALIGN.CENTER)
    create_textbox(slide, Inches(7.3), Inches(3.8), Inches(4.5), Inches(1.0), "$10.2 Million", 44, RGBColor(100, 255, 100), bold=True, align=PP_ALIGN.CENTER)
    create_textbox(slide, Inches(7.3), Inches(4.8), Inches(4.5), Inches(0.5), "Assuming 20% Intervention Mitigation", 14, TEXT_SUB, align=PP_ALIGN.CENTER)


    prs.save('Workflow_Presentation.pptx')
    print("Professional Presentation generated successfully!")

if __name__ == '__main__':
    create_presentation()
